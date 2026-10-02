"""Extract plain text from raw/ sources into _meta/extracted/ (a disposable cache).

raw/ is never modified. Re-run any time new sources are added:
    .venv/bin/python tools/extract.py            # only new/changed files
    .venv/bin/python tools/extract.py --force    # re-extract everything
"""
import re
import sys
from pathlib import Path

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "raw"
OUT = ROOT / "_meta" / "extracted"
TEXT_EXT = {".py", ".pl", ".md", ".txt"}


def pptx_text(path):
    parts = []
    for i, slide in enumerate(Presentation(path).slides, 1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines += [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    lines.append(" | ".join(c.text for c in row.cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"[notes] {notes}")
        parts.append(f"--- slide {i} ---\n" + "\n".join(lines))
    return "\n\n".join(parts)


def pdf_text(path):
    reader = PdfReader(path)
    return "\n\n".join(
        f"--- page {i} ---\n{page.extract_text() or ''}" for i, page in enumerate(reader.pages, 1)
    )


def docx_text(path):
    doc = Document(path)
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text for c in row.cells))
    return "\n".join(lines)


HANDLERS = {".pptx": pptx_text, ".pdf": pdf_text, ".docx": docx_text}


def main(force=False):
    for src in sorted(RAW.rglob("*")):
        if not src.is_file() or src.name.startswith(("~$", ".")) or "__pycache__" in src.parts:
            continue
        ext = src.suffix.lower()
        if ext not in HANDLERS and ext not in TEXT_EXT:
            continue
        dst = OUT / src.relative_to(RAW).with_suffix(src.suffix + ".txt")
        if not force and dst.exists() and dst.stat().st_mtime >= src.stat().st_mtime:
            continue
        dst.parent.mkdir(parents=True, exist_ok=True)
        try:
            text = HANDLERS[ext](src) if ext in HANDLERS else src.read_text(errors="replace")
        except Exception as e:
            print(f"FAILED {src.relative_to(RAW)}: {e}")
            continue
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
        dst.write_text(text)
        print(f"{len(text):>9} chars  {src.relative_to(RAW)}")


if __name__ == "__main__":
    main(force="--force" in sys.argv)
